# -*- coding: utf-8 -*-
"""整合版全组 PPT 优化：删页码、统一主标题、底部导航栏（当前部分高亮）"""
import re, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = "merged.pptx"
OUT = "FF5全组汇报_导航版.pptx"
TEAL = RGBColor(0x01, 0xA5, 0xB1)
DARK = RGBColor(0x1F, 0x1F, 0x1F)
GREY = RGBColor(0x8A, 0x8A, 0x8A)
LINE = RGBColor(0xDD, 0xDD, 0xDD)

TABS = ["开场与检索", "引言与文献回顾", "研究方法", "适用性检验", "分时期检验", "动量因子", "结论与反思", "方法论与SOP"]
# 每个部分覆盖的页码（1-based，含端点）
GROUPS = [(2, 2), (3, 12), (13, 18), (19, 24), (25, 34), (35, 39), (40, 41), (42, 44)]

PAGE_RE = re.compile(r"^\d{1,2}\s*/\s*\d{1,3}$")
BAR_Y = 7.18  # in

def set_font(run, size, bold, color, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)

def no_line(shape):
    shape.line.fill.background()

prs = Presentation(SRC)
slides = list(prs.slides)
n = len(slides)

def group_of(idx1):  # idx1: 1-based
    for gi, (a, b) in enumerate(GROUPS):
        if a <= idx1 <= b:
            return gi
    return None

for i, slide in enumerate(slides, 1):
    shapes = list(slide.shapes)
    # ---- 1) 删页码 + 旧页脚水印 ----
    for sh in shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if txt and (PAGE_RE.match(txt) or "三维结构法" in txt):
            sh._element.getparent().remove(sh._element)

    # ---- 1b) 防遮挡：正文元素 bottom 超过导航栏顶线则上移 ----
    shapes2 = list(slide.shapes)
    over = []
    for sh in shapes2:
        if sh.top is None or sh.height is None:
            continue
        if sh.height / 914400 >= 6.8:      # 全页背景不动
            continue
        if (sh.top + sh.height) / 914400 > BAR_Y - 0.05:
            over.append((sh, (sh.top + sh.height) / 914400))
    if i == 34:
        # P34 左大图+右侧卡片组为整体：除背景/页脚外全部内容元素统一上移
        for sh in shapes2:
            if sh.top is None or sh.height is None:
                continue
            if sh.height / 914400 >= 6.8:
                continue
            if sh.has_text_frame and (PAGE_RE.match(sh.text_frame.text.strip()) or "三维结构法" in sh.text_frame.text):
                continue
            sh.top = Emu(int(sh.top - 0.14 * 914400))
    else:
        # 每个越界元素各自对齐到 7.12in（底部留 0.06in 余量）
        for sh, bottom_in in over:
            sh.top = Emu(int(sh.top - (bottom_in - 7.12) * 914400))

    # ---- 2) 统一内容页主标题（跳过封面；扉页无顶部大字不受影响）----
    if i != 1:
        cands = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            top = sh.top if sh.top is not None else 0
            if top > Emu(int(1.25 * 914400)):
                continue
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if not t or PAGE_RE.match(t):
                    continue
                sz = next((r.font.size.pt for r in para.runs if r.font.size), None)
                if sz and sz >= 18:
                    cands.append((top, sh, sz, t))
                    break
        if cands:
            cands.sort(key=lambda x: x[0])
            _, sh, _, _ = cands[0]
            sh.left, sh.top = Inches(0.58), Inches(0.28)
            sh.width, sh.height = Inches(11.6), Inches(0.6)
            tf = sh.text_frame
            tf.word_wrap = True
            for para in tf.paragraphs:
                if not para.runs:
                    continue
                for r_ in para.runs:
                    set_font(r_, 22, True, DARK)

    # ---- 3) 底部导航栏（封面不放）----
    if i == 1:
        continue
    gi = group_of(i)
    BAR_H = Inches(0.32)
    bar_y = prs.slide_height - BAR_H
    # 背景白条 + 顶部细线
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_y, prs.slide_width, BAR_H)
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    no_line(bar)
    bar.shadow.inherit = False
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_y, prs.slide_width, Emu(9525))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE
    no_line(ln)
    ln.shadow.inherit = False
    # tabs
    tw = prs.slide_width / len(TABS)
    for ti, label in enumerate(TABS):
        active = (ti == gi)
        tx = tw * ti
        if active:
            blk = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, bar_y, tw, BAR_H)
            blk.fill.solid(); blk.fill.fore_color.rgb = TEAL
            no_line(blk)
            blk.shadow.inherit = False
        tb = slide.shapes.add_textbox(tx + Inches(0.05), bar_y, tw - Inches(0.10), BAR_H)
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = label
        set_font(run, 11, active, RGBColor(0xFF, 0xFF, 0xFF) if active else GREY)

# ================= 追加：07 方法论提炼与论文 SOP（P42-44） =================
DECO = r"C:/Users/liuko/WorkBuddy/2026-09-13-07-29-45/FF5第一二章精读/assets/deco_teal.png"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF5, 0xFB, 0xFB)
GOLD_BG = RGBColor(0xFF, 0xF8, 0xEE)
GOLD_TX = RGBColor(0x9A, 0x6A, 0x12)
BODY = RGBColor(0x44, 0x44, 0x44)
SUB = RGBColor(0x77, 0x77, 0x77)

def add_rect(s, x, y, w, h, fill=None, line=None, lw=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw or 1)
    sp.shadow.inherit = False
    return sp

def add_par(s, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """lines: list of (text, size, bold, color)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    first = True
    for text, size, bold, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = 1.25
        r = p.add_run(); r.text = text
        set_font(r, size, bold, color)
    return tb

def draw_nav(s, gi):
    BAR_H = Inches(0.32)
    bar_y = prs.slide_height - BAR_H
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_y, prs.slide_width, BAR_H)
    bar.fill.solid(); bar.fill.fore_color.rgb = WHITE
    no_line(bar); bar.shadow.inherit = False
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_y, prs.slide_width, Emu(9525))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE
    no_line(ln); ln.shadow.inherit = False
    tw = prs.slide_width / len(TABS)
    for ti, label in enumerate(TABS):
        active = (ti == gi)
        tx = tw * ti
        if active:
            blk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, tx, bar_y, tw, BAR_H)
            blk.fill.solid(); blk.fill.fore_color.rgb = TEAL
            no_line(blk); blk.shadow.inherit = False
        tb = s.shapes.add_textbox(tx + Inches(0.05), bar_y, tw - Inches(0.10), BAR_H)
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        set_font(r, 11, active, WHITE if active else GREY)

def content_header(s, title):
    add_rect(s, 0, 0, 13.333, 0.06, fill=TEAL)
    add_par(s, 0.58, 0.28, 11.6, 0.6, [(title, 22, True, DARK)])

blank = None
for ly in prs.slide_layouts:
    if not ly.placeholders or len(ly.placeholders) == 0:
        blank = ly; break
if blank is None:
    blank = prs.slide_layouts[6]

# ---- P42 扉页 07 ----
s42 = prs.slides.add_slide(blank)
s42.shapes.add_picture(DECO, 0, 0, width=prs.slide_width, height=prs.slide_height)
box = add_rect(s42, 1.27, 2.56, 2.17, 2.17, fill=None, line=WHITE, lw=3)
add_par(s42, 1.27, 2.56, 2.17, 2.17, [("07", 54, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
add_par(s42, 4.54, 2.98, 7.9, 0.9, [("方法论提炼与论文 SOP", 34, True, DARK)])
add_rect(s42, 4.54, 3.98, 0.92, 0.06, fill=TEAL)
add_par(s42, 4.54, 4.28, 7.7, 1.3, [
    ("把文献拆解的产出沉淀为团队资产：三维结构法逆向拆解", 15, False, RGBColor(0x5A, 0x5A, 0x5A)),
    ("的方法论，以及从检索到归档的六步论文 SOP 标准操作流程。", 15, False, RGBColor(0x5A, 0x5A, 0x5A)),
])
draw_nav(s42, 7)

# ---- P43 方法论：三维结构法 ----
s43 = prs.slides.add_slide(blank)
content_header(s43, "7.1 团队方法论：三维结构法逆向拆解")
STEPS = [
    ("步骤一", "Trace the line · 摸清学术家谱",
     "不读公式先拉时间轴：到 GitHub / Papers with Code 把细分赛道近 3 年的进化链条拉出来，看原作者遇到了什么瓶颈、后来者用什么组件查漏补缺——看懂进化套路，选题就成功了一半。",
     "本篇实操：画出母模型主线＋本土支线双行家谱，锁定三条裂缝——体系内 CMA 被判冗余、同市场三种相反结论、体系外 LSY 指出壳污染与 E/P 口径。"),
    ("步骤二", "Extract narrative · 提取高阶语料",
     "高级转折句、精妙因果推导、把魔改代码包装成经典数学框架的描述，直接标黄摘录——这就是未来的专属语料库，写论文时不必憋初级简单句。",
     "本篇实操：摘录五类语料，含 P205 结论段「某些分组下仍无法通过 GRS 检验」的自我设限句——作者亲手留下的空白任务单。"),
    ("步骤三", "Run the code · 跑通灵魂代码",
     "克隆开源仓库、逐行精简写注释：公式与代码之间的实现 Gap，往往就是发高分 SCI 的最佳创新点——这是最关键、最拉开差距的一步。",
     "本篇实操：复现骨架跑通「2×3 分组 → GRS 检验 → 邹至庄断点检验」全链路，产出 8 条可立项的 Gap 清单。"),
]
y = 0.98
for tag, en, body, act in STEPS:
    add_rect(s43, 0.56, y, 12.22, 1.62, fill=CARD, line=RGBColor(0xDD, 0xEE, 0xEE), lw=1)
    add_rect(s43, 0.56, y, 1.30, 1.62, fill=TEAL)
    add_par(s43, 0.60, y, 1.22, 1.62, [(tag, 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    add_par(s43, 2.02, y + 0.10, 10.55, 1.45, [
        (en, 14.5, True, RGBColor(0x0B, 0x7C, 0x86)),
        (body, 12.5, False, BODY),
        (act, 12, False, RGBColor(0x77, 0x77, 0x77)),
    ])
    y += 1.74
add_rect(s43, 0.56, 6.22, 12.22, 0.72, fill=GOLD_BG)
add_rect(s43, 0.56, 6.22, 0.06, 0.72, fill=RGBColor(0xD8, 0x95, 0x2F))
add_par(s43, 0.82, 6.22, 11.8, 0.72,
        [("读文献的产出不是一个「理解」，而是一个「选题」——三维结构法把每一篇论文都变成下一步研究的入口。", 14, True, GOLD_TX)],
        anchor=MSO_ANCHOR.MIDDLE)
draw_nav(s43, 7)

# ---- P44 论文 SOP 六步 ----
s44 = prs.slides.add_slide(blank)
content_header(s44, "7.2 论文 SOP：从检索到归档的六步标准操作流程")
SOPS = [
    ("01", "检索定篇", "知网专业检索 LY='金融研究' AND SU='因子'，共 105 条；按被引量降序锁定精读对象（本文被引 1036、下载 27287）。"),
    ("02", "三维拆解", "家谱 → 语料 → 代码三步走，产出三份中间产物：裂缝清单、专属语料库、Gap 清单。"),
    ("03", "选题立项", "从 Gap 中筛「前人没做但可证伪」的改写方向：如壳污染修正口径、E/P 替代 B/M、未知断点扫描。"),
    ("04", "实验设计", "设对照实验与稳健性检验；所有口径标注原文页码与表号，杜绝「想当然」。"),
    ("05", "写作成稿", "用语料库升级表达，方法—数据—结论逐环对应；讲稿与实际能力对齐，不夸大。"),
    ("06", "归档复用", "演示稿、复现脚本、运行日志、语料库、Gap 清单打包上传 GitHub 仓库；SOP 随每次精读迭代更新。"),
]
pos = [(0.56, 1.00), (4.85, 1.00), (9.14, 1.00), (0.56, 3.28), (4.85, 3.28), (9.14, 3.28)]
for (num, title, body), (x, y) in zip(SOPS, pos):
    add_rect(s44, x, y, 3.98, 2.10, fill=WHITE, line=RGBColor(0xE2, 0xE2, 0xE2), lw=1)
    add_rect(s44, x, y, 3.98, 0.10, fill=TEAL)
    add_par(s44, x + 0.16, y + 0.22, 3.66, 1.80, [
        (num + "　" + title, 15, True, RGBColor(0x0B, 0x7C, 0x86)),
        (body, 12, False, BODY),
    ])
add_rect(s44, 0.56, 5.62, 12.22, 1.28, fill=GOLD_BG)
add_rect(s44, 0.56, 5.62, 0.06, 1.28, fill=RGBColor(0xD8, 0x95, 0x2F))
add_par(s44, 0.82, 5.62, 11.8, 1.28, [
    ("◆ 收尾呼应开场：本次分享的全部交付物已按第 06 步归档至团队 GitHub 项目仓库", 14.5, True, GOLD_TX),
    ("文献解读的终点不是一份 PPT，而是一套可复用的 SOP——下一篇论文，按流程再跑一遍。", 13, False, RGBColor(0x4A, 0x4A, 0x4A)),
], anchor=MSO_ANCHOR.MIDDLE)
draw_nav(s44, 7)

prs.save(OUT)
print("saved", OUT, "slides:", n)